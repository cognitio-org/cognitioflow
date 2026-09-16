"""Tests: Phase 2 — every file and audio route goes through storage.py."""
import base64
import io
import pathlib

import pytest

import storage

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def _docx(text: str) -> bytes:
    import docx
    d = docx.Document(); d.add_paragraph(text); b = io.BytesIO(); d.save(b)
    return b.getvalue()


def _pptx(text: str) -> bytes:
    from pptx import Presentation
    p = Presentation(); p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = text; b = io.BytesIO(); p.save(b)
    return b.getvalue()


def _cid(client):
    return client.get("/api/courses").json()[0]["id"]


def _upload(client, cid, name, data, ctype="application/octet-stream"):
    r = client.post(f"/api/courses/{cid}/files", files={"file": (name, io.BytesIO(data), ctype)}, data={"week": "1"})
    assert r.status_code == 200, r.text
    return r.json()


def _key(pg, table, id_):
    return pg.execute(f"SELECT key FROM {table} WHERE id=%s", (id_,)).fetchone()[0]


@pytest.mark.parametrize("name,data,status,needle", [
    ("Week 1 slides.pdf", _pdf("Article 34 TFEU"), "indexed", "Article 34"),
    ("Week 1 lecture.pptx", _pptx("Dassonville formula"), "indexed", "Dassonville"),
    ("WG notes.docx", _docx("Cassis de Dijon"), "indexed", "Cassis"),
    ("transcript.md", "Keck — selling arrangements".encode(), "indexed", "Keck"),
    ("whiteboard.png", PNG, "image", None),
])
def test_upload_stores_bytes_under_key_and_indexes_text(client, pg, name, data, status, needle):
    cid = _cid(client)
    body = _upload(client, cid, name, data)
    assert body["status"] == status
    key = _key(pg, "files", body["id"])
    assert key == f"courses/{cid}/files/{body['id']}/{name}"
    assert storage.get(key) == data
    if needle:
        assert needle in client.get(f"/api/files/{body['id']}/text").json()["text"]


def test_raw_download_streams_original_bytes_with_filename(client):
    data = _pdf("Schütze chapter 3")
    fid = _upload(client, _cid(client), "Schütze ch3.pdf", data)["id"]
    r = client.get(f"/api/files/{fid}/raw")
    assert r.status_code == 200
    assert r.content == data
    assert r.headers["content-type"] == "application/pdf"
    assert "filename*=UTF-8''Sch%C3%BCtze%20ch3.pdf" in r.headers["content-disposition"]


def test_raw_download_redirects_when_storage_can_sign(client, monkeypatch):
    fid = _upload(client, _cid(client), "a.pdf", b"%PDF")["id"]
    monkeypatch.setattr(storage, "url", lambda key, *a, **k: f"https://signed.example/{key}?X-Goog-Signature=abc")
    r = client.get(f"/api/files/{fid}/raw", follow_redirects=False)
    assert r.status_code == 302
    assert r.headers["location"].startswith("https://signed.example/courses/")


def test_delete_file_removes_object_and_row(client, pg):
    fid = _upload(client, _cid(client), "gone.txt", b"bye")["id"]
    key = _key(pg, "files", fid)
    assert client.delete(f"/api/files/{fid}").status_code == 200
    assert not storage.exists(key)
    assert pg.execute("SELECT COUNT(*) FROM files WHERE id=%s", (fid,)).fetchone()[0] == 0


def test_missing_or_legacy_object_is_404_not_500(client, pg):
    cid = _cid(client)
    a = _upload(client, cid, "a.pdf", b"%PDF-a")["id"]
    storage.delete(_key(pg, "files", a))  # object vanished behind the app's back
    assert client.get(f"/api/files/{a}/raw").status_code == 404
    b = _upload(client, cid, "b.pdf", b"%PDF-b")["id"]
    pg.execute("UPDATE files SET key=%s WHERE id=%s", ("/Users/someone/Desktop/cognitioflow/data/uploads/b.pdf", b)); pg.commit()
    assert client.get(f"/api/files/{b}/raw").status_code == 404


def test_tutor_context_reads_images_from_storage_and_caps_them(client):
    from run import MAX_IMAGES, build_context
    cid = _cid(client)
    for i in range(MAX_IMAGES + 1):
        _upload(client, cid, f"board{i}.png", PNG, "image/png")
    _, images, _ = build_context(cid)
    assert len(images) == MAX_IMAGES
    assert images[0]["source"] == {"type": "base64", "media_type": "image/png", "data": base64.b64encode(PNG).decode()}


def test_tutor_context_skips_image_missing_from_storage(client, pg):
    from run import build_context
    cid = _cid(client)
    fid = _upload(client, cid, "lost.png", PNG, "image/png")["id"]
    storage.delete(_key(pg, "files", fid))
    assert build_context(cid)[1] == []


def test_note_to_file_writes_markdown_object(client, pg):
    cid = _cid(client)
    nid = client.post(f"/api/courses/{cid}/notes", json={"title": "Free movement", "body": "# Goods\nArt 34 — Dassonville"}).json()["id"]
    fid = client.post(f"/api/notes/{nid}/to-file").json()["id"]
    key = _key(pg, "files", fid)
    assert key == f"courses/{cid}/files/{fid}/Free movement.md"
    assert storage.get(key).decode() == "# Goods\nArt 34 — Dassonville"
    r = client.get(f"/api/files/{fid}/raw")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/markdown")


def test_recording_roundtrip_with_range_requests(client, pg):
    cid = _cid(client)
    nid = client.post(f"/api/courses/{cid}/notes", json={"title": "Lecture 3"}).json()["id"]
    rid = client.post(f"/api/notes/{nid}/recordings/start").json()["id"]
    audio = bytes(range(256)) * 20  # 5120 bytes
    r = client.post(f"/api/recordings/{rid}/finish", files={"audio": ("r.webm", io.BytesIO(audio), "audio/webm")},
                    data={"seconds": "12.5", "auto": "0"})
    assert r.status_code == 200, r.text
    key = _key(pg, "recordings", rid)
    assert key == f"notes/{nid}/audio/{rid}.webm"
    assert [x["id"] for x in client.get(f"/api/notes/{nid}/recordings").json()] == [rid]

    full = client.get(f"/api/recordings/{rid}/audio")
    assert full.status_code == 200 and full.content == audio
    assert full.headers["accept-ranges"] == "bytes" and full.headers["content-type"] == "audio/webm"

    part = client.get(f"/api/recordings/{rid}/audio", headers={"Range": "bytes=100-199"})
    assert part.status_code == 206
    assert part.headers["content-range"] == f"bytes 100-199/{len(audio)}"
    assert part.content == audio[100:200]
    assert client.get(f"/api/recordings/{rid}/audio", headers={"Range": "bytes=5110-"}).content == audio[5110:]
    assert client.get(f"/api/recordings/{rid}/audio", headers={"Range": "bytes=-10"}).content == audio[-10:]
    assert client.get(f"/api/recordings/{rid}/audio", headers={"Range": "bytes=9000-"}).status_code == 416

    assert client.delete(f"/api/recordings/{rid}").status_code == 200
    assert not storage.exists(key)
    assert client.get(f"/api/recordings/{rid}/audio").status_code == 404


def test_finish_unknown_recording_is_404(client):
    r = client.post("/api/recordings/nope/finish", files={"audio": ("r.webm", io.BytesIO(b"x"), "audio/webm")}, data={"auto": "0"})
    assert r.status_code == 404


def test_watched_folder_is_gone(client):
    cid = _cid(client)
    assert "watch_root" not in client.get("/api/config").json()
    assert client.get(f"/api/courses/{cid}/watch").status_code in (404, 405)
    assert client.post(f"/api/courses/{cid}/scan").status_code in (404, 405)


def test_run_py_has_no_filesystem_access_to_user_content():
    src = (pathlib.Path(__file__).parent.parent / "run.py").read_text()
    for banned in ("UPLOADS", "AUDIO", "WATCH", "read_bytes", "write_bytes", "google.cloud"):
        assert banned not in src, banned
    lines = [l for l in src.splitlines() if "FileResponse" in l and not l.startswith("from fastapi")]
    assert len(lines) == 1 and "index.html" in lines[0]
